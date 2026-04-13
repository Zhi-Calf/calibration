
from pathlib import Path
import glob, os
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ORIGIN_DIR = Path(r"C:\\Users\\Xiao\\Downloads\\val_origin_results")
NOISE_DIR = Path(r"C:\\Users\\Xiao\\Downloads\\val_results")
OUT = NOISE_DIR / "eval_template_filled.xlsx"

MODELS=['yolov13','detr','co_detr','vitdet','diffusiondet']
SCENS=['photoelectric_signal','optical_distortion','radiometric','Environmental-Interference','Object-level-noise-category','Environmental-Coupling','Time-and-space-mismatch','Cross-sensor-interference']

def to_percent(v): return v*100.0 if v<=1.0 else v

def read_text(p):
    if p is None or not p.exists(): return ''
    return p.read_text(encoding='utf-8',errors='ignore')

def origin_file(model, modality):
    if model=='yolov13':
        name='yolov13_singlemodal_data_class15_train2_val.md' if modality=='singlemodal' else 'yolov13_multimodal_multidata_rgbd_val.md'
        p=ORIGIN_DIR/name; return p if p.exists() else None
    p=ORIGIN_DIR/f'{model}_{modality}_origin_val.md'; return p if p.exists() else None

def noise_files(model, modality):
    if modality=='singlemodal':
        p_no=NOISE_DIR/f'{model}_noise_val_metrics.md'; p_yes=NOISE_DIR/f'{model}_noise_val_metrics_singlemodal_noise.md'
    else:
        p_no=NOISE_DIR/f'{model}_multimodal_rgbd_val_original.md'; p_yes=NOISE_DIR/f'{model}_multimodal_rgbd_val_noise.md'
    return (p_no if p_no.exists() else None, p_yes if p_yes.exists() else None)

def parse_origin_pr(text):
    lines=text.splitlines()
    for i,line in enumerate(lines):
        if 'mAP50-95' in line and line.strip().startswith('|'):
            for j in range(i+1,min(i+6,len(lines))):
                row=lines[j].strip();
                if not row.startswith('|'): continue
                cols=[c.strip() for c in row.strip('|').split('|')]
                if len(cols)<7: continue
                try: return {'P':to_percent(float(cols[2])),'R':to_percent(float(cols[3]))}
                except: pass
    if 'P (%)' in text and 'R (%)' in text:
        for row in lines:
            row=row.strip();
            if not row.startswith('|'): continue
            cols=[c.strip() for c in row.strip('|').split('|')]
            if len(cols)<9: continue
            try: return {'P':to_percent(float(cols[7])),'R':to_percent(float(cols[8]))}
            except: pass
    return {}

def parse_noise(text):
    out={}
    for line in text.splitlines():
        row=line.strip()
        if not (row.startswith('| scenario_val_') and row.endswith('|')): continue
        cols=[c.strip() for c in row.strip('|').split('|')]
        if len(cols)<8: continue
        try: out[cols[0]]={'P':float(cols[6]),'R':float(cols[7])}
        except: pass
    return out

cands=sorted(glob.glob(r'C:\\Users\\Xiao\\Downloads\\val_results\\*.xlsx'))
TEMPLATE=None
for f in cands:
    b=os.path.basename(f)
    if 'expanded' in b.lower() or '_???' in b or 'eval_template_filled' in b:
        continue
    TEMPLATE=Path(f); break

metrics={}
for modality in ['singlemodal','multimodal']:
    for model in MODELS:
        ori=parse_origin_pr(read_text(origin_file(model,modality)))
        fno,fyes=noise_files(model,modality)
        metrics[(modality,model)]={'origin':ori,'no':parse_noise(read_text(fno)),'yes':parse_noise(read_text(fyes))}

wb=load_workbook(TEMPLATE)
ws=wb[wb.sheetnames[0]]
for m in list(ws.merged_cells.ranges):
    try: ws.unmerge_cells(str(m))
    except: pass
for col in [12,11,10,9,8,7,6,5,3]: ws.insert_cols(col+1,1)
for r in range(1,5):
    for c in range(1,26): ws.cell(r,c).value=None

ws.cell(1,1,'Data'); ws.cell(2,1,'Modal')
ws.cell(1,2,'Model'); ws.cell(1,3,'Origin'); ws.cell(1,5,'Finetune')
ws.cell(1,6,'RGB'); ws.cell(1,12,'Depth'); ws.cell(1,16,'Coupling'); ws.cell(1,22,'Overall')
for c,name in [(6,'Photoelectric signal'),(8,'Optical distortion'),(10,'Radiometric measurement'),(12,'Environmental interference'),(14,'Object-level'),(16,'Environmental Coupling'),(18,'Time and Space'),(20,'Cross_sensor interference'),(22,'NO'),(24,'YES')]:
    ws.cell(2,c,name); ws.cell(4,c,'P'); ws.cell(4,c+1,'R')

for rg in ['A1:A4','B1:B4','C1:D4','E1:E4','F1:K1','L1:O1','P1:U1','V1:Y1','F2:G3','H2:I3','J2:K3','L2:M3','N2:O3','P2:Q3','R2:S3','T2:U3','V2:W3','X2:Y3','A5:A14','A15:A24','B5:B6','B7:B8','B9:B10','B11:B12','B13:B14','B15:B16','B17:B18','B19:B20','B21:B22','B23:B24','C5:C6','C7:C8','C9:C10','C11:C12','C13:C14','C15:C16','C17:C18','C19:C20','C21:C22','C23:C24']:
    ws.merge_cells(rg)

ws.cell(5,1,'Single-modal Methods'); ws.cell(15,1,'Multi-modal Methods')
for i,mn in enumerate(['Yolo13','DETR','CO-DETR','VITDET','DiffusionDet']):
    r=5+2*i
    ws.cell(r,2,mn); ws.cell(r,5,'NO'); ws.cell(r+1,5,'YES')

col_origin=(3,4)
scen_cols={'photoelectric_signal':(6,7),'optical_distortion':(8,9),'radiometric':(10,11),'Environmental-Interference':(12,13),'Object-level-noise-category':(14,15),'Environmental-Coupling':(16,17),'Time-and-space-mismatch':(18,19),'Cross-sensor-interference':(20,21)}
col_overall_no=(22,23); col_overall_yes=(24,25)

for start, modality in [(5,'singlemodal'),(15,'multimodal')]:
    for i,model in enumerate(MODELS):
        no_row=start+2*i; yes_row=no_row+1; d=metrics[(modality,model)]
        if 'P' in d['origin']: ws.cell(no_row,col_origin[0],round(d['origin']['P'],4))
        if 'R' in d['origin']: ws.cell(no_row,col_origin[1],round(d['origin']['R'],4))
        all_no=d['no'].get('scenario_val_all',{})
        if 'P' in all_no: ws.cell(no_row,col_overall_no[0],round(all_no['P'],4))
        if 'R' in all_no: ws.cell(no_row,col_overall_no[1],round(all_no['R'],4))
        all_yes=d['yes'].get('scenario_val_all',{})
        if 'P' in all_yes: ws.cell(yes_row,col_overall_yes[0],round(all_yes['P'],4))
        if 'R' in all_yes: ws.cell(yes_row,col_overall_yes[1],round(all_yes['R'],4))
        for s in SCENS:
            key=f'scenario_val_{s}'; c1,c2=scen_cols[s]
            no=d['no'].get(key,{}); yes=d['yes'].get(key,{})
            if 'P' in no: ws.cell(no_row,c1,round(no['P'],4))
            if 'R' in no: ws.cell(no_row,c2,round(no['R'],4))
            if 'P' in yes: ws.cell(yes_row,c1,round(yes['P'],4))
            if 'R' in yes: ws.cell(yes_row,c2,round(yes['R'],4))

center=Alignment(horizontal='center',vertical='center',wrap_text=True)
for r in range(1,25):
    for c in range(1,26): ws.cell(r,c).alignment=center
for r in range(1,5):
    for c in range(1,26): ws.cell(r,c).font=Font(name='Calibri',size=11,bold=True)
for c,w in {1:20,2:14,3:9,4:9,5:9}.items(): ws.column_dimensions[get_column_letter(c)].width=w
for c in range(6,26): ws.column_dimensions[get_column_letter(c)].width=10.5

wb.save(OUT)
print('saved',OUT)
